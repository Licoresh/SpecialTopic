from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from src.reporting import export_csv, export_excel, export_html
from src.analytics import filter_frame
from src.visualization import bar, heatmap, distribution

def test_exports_formula_safety_and_html():
    frame = pd.DataFrame({"value":["=1+1", "<script>bad</script>", "SQL"],"count":[1,2,3]})
    workbook = load_workbook(BytesIO(export_excel({"test":frame}, {"seed":42})))
    assert workbook.sheetnames == ["test", "analysis_metadata"]
    assert workbook["test"]["A2"].data_type == "s"
    assert b"'=1+1" in export_csv(frame)
    assert b"&lt;script&gt;" in export_html({"test":frame}, {})

def test_filters_and_chart_serialization():
    frame = pd.DataFrame({"category":["A","A","B"], "count":[1,2,3]})
    filtered = filter_frame(frame, {"category":["A"]})
    assert len(filtered) == 2
    assert len(bar(filtered, "category", "count", "Test").data[0].x) == 2
    assert distribution(filtered, "category", "Test").to_json()
    assert heatmap(pd.DataFrame([[1,2],[3,4]]), "Test").to_json()

def test_workbook_preserves_long_job_description_and_metadata():
    text = "Long job description. " * 3000
    frame = pd.DataFrame({"job_description":[text], "control":["text\x01value"]})
    book = load_workbook(BytesIO(export_excel({"jobs":frame}, {"long_setting":text})))
    parts = list(book["long_text"].values)[1:]
    assert "".join(row[4] for row in parts) == text
    assert book["jobs"]["B2"].value == "text\\u0001value"
    metadata = list(book["analysis_metadata"].values)[1:]
    import json
    assert json.loads("".join(row[2] for row in metadata)) == text
