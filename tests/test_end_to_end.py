from io import BytesIO
from openpyxl import load_workbook
import pandas as pd
from src import storage
from src.classification import compare_models, save_model
from src.labeling import save_annotations
from src.workflow import run_extraction, analysis_bundle, input_signature, current_extractions
from src.reporting import export_excel
import pytest

def test_input_to_workbook(workspace, jobs, taxonomy, labeled):
    for kind, frame in (("jobs",jobs),("psf_aai",taxonomy)):
        storage.save_frame("cleaned_"+kind, frame)
    save_annotations(labeled, "SYNTHETIC TEST REVIEWER")
    result = compare_models(storage.read_frame("labeled_competencies"))["Logistic Regression"]
    save_model(result["model"], workspace/"models/competency_classifier.joblib")
    storage.save_json("selected_model", {"name":"Logistic Regression", "metrics":result["metrics"]})
    extracted = run_extraction()
    assert extracted.ml_category.ne("").any()
    assert extracted.psf_competency.ne("").any()
    tables, meta = analysis_bundle()
    assert not tables["employer_skill_demand"].empty
    assert {"psf_mapping", "alignment_results", "unmapped_skills", "emerging_skills", "verified_psf_aai", "extracted_job_skills"} <= set(tables)
    assert meta["extraction"]["signature"] == input_signature()
    snapshot = {"jobs": storage.read_frame("cleaned_jobs"),
                "taxonomy": storage.read_frame("cleaned_psf_aai"),
                "extracted": current_extractions()}
    cached_path_tables, _ = analysis_bundle(source_snapshot=snapshot)
    pd.testing.assert_frame_equal(tables["employer_skill_demand"], cached_path_tables["employer_skill_demand"])
    book = load_workbook(BytesIO(export_excel(tables, meta)))
    assert "employer_skill_demand" in book.sheetnames
    assert book["employer_skill_demand"].max_row > 1
    filtered, _ = analysis_bundle(job_filters={"job_title":["Synthetic developer"]})
    assert len(filtered["cleaned_jobs"]) == 1
    assert filtered["employer_skill_demand"].employer_demand_pct.eq(100).all()

def test_changed_upload_requires_cleaning(workspace, jobs):
    storage.save_frame("cleaned_jobs", jobs)
    before = input_signature()
    storage.archive_raw(b"description\nnew data", "jobs.csv", "jobs")
    assert input_signature() != before
    with pytest.raises(ValueError, match="mapped and cleaned"):
        run_extraction()


def test_unreviewed_job_provenance_blocks_extraction(workspace, jobs):
    from src.workflow import run_extraction
    from src.utils import frame_hash
    storage.archive_raw(b"job_title,job_description\\nSynthetic,SQL", "jobs.csv", "jobs")
    upload = storage.read_json("uploads")["jobs"]
    jobs["_dataset_hash"] = upload["hash"]
    storage.save_frame("cleaned_jobs", jobs)
    storage.save_frame("mapped_jobs", jobs)
    storage.save_json("cleaning_jobs", {"source_hash": frame_hash(jobs)})
    with pytest.raises(ValueError, match="Philippine provenance"):
        run_extraction()


def test_research_exports_exclude_external_training_sources(workspace, jobs, taxonomy):
    storage.save_frame("cleaned_jobs", jobs)
    storage.save_frame("cleaned_psf_aai", taxonomy)
    selected_id = jobs.source_id.iloc[0]
    storage.save_frame("labeled_competencies", pd.DataFrame([
        {"phrase": "SQL", "normalized_phrase": "sql", "label": "Data / Analytics", "source_type": "jobs", "source_id": selected_id},
        {"phrase": "external skill", "normalized_phrase": "external skill", "label": "Technical / Digital", "source_type": "jobs", "source_id": "EXTERNAL-TRAINING-ONLY"},
    ]))
    tables, metadata = analysis_bundle(source_snapshot={"jobs": jobs.iloc[:1], "taxonomy": taxonomy, "extracted": pd.DataFrame()})
    assert tables["labeled_competencies"].source_id.tolist() == [selected_id]
    assert len(storage.read_frame("labeled_competencies")) == 2
    assert "test_predictions" not in tables
    assert "model_metrics" not in tables
    assert metadata["research_sample"]["source_ids"] == [selected_id]
    book = load_workbook(BytesIO(export_excel(tables, metadata)))
    assert "test_predictions" not in book.sheetnames
    assert book["labeled_competencies"].max_row == 2
